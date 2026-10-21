"""Tests for the slidestyle builder, the template and the checker.

Run with:
    uvx --with python-pptx --with lxml --with pillow --with pypandoc_binary \
        pytest skills/slides-style/tests
"""

import importlib.util
import pathlib
import shutil
import subprocess
import sys
import types
import zipfile

from lxml import etree
from PIL import Image
import pptx
import pptx.enum.shapes
from pptx.oxml import ns
import pytest

_ROOT = pathlib.Path(__file__).parents[1]


def _load(name: str, path: pathlib.Path) -> types.ModuleType:
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise ImportError(f"cannot load {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


slidestyle = _load("slidestyle", _ROOT / "assets" / "slidestyle.py")
checker = _load("check_slides", _ROOT / "scripts" / "check_slides.py")
svg_tool = _load("latex_to_svg", _ROOT / "scripts" / "latex_to_svg.py")

_HAS_PANDOC = bool(shutil.which("pandoc")) or bool(
    importlib.util.find_spec("pypandoc")
)
_HAS_LATEX = bool(shutil.which("pdflatex") and shutil.which("pdftocairo"))
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"


@pytest.fixture(autouse=True)
def _no_local_assets(tmp_path, monkeypatch):
    monkeypatch.setenv(slidestyle.ASSETS_ENV, str(tmp_path / "no-assets"))


@pytest.fixture
def png(tmp_path):
    path = tmp_path / "figure.png"
    Image.new("RGB", (800, 400), "white").save(path)
    return path


def _codes(findings):
    return {finding.code for finding in findings if finding.severity == "ERROR"}


def test_template_is_widescreen_with_house_layouts():
    deck = slidestyle.new_deck()
    assert (deck.slide_width, deck.slide_height) == (12192000, 6858000)
    assert [layout.name for layout in deck.slide_layouts] == [
        "Title",
        "Content",
        "Blank",
    ]
    with zipfile.ZipFile(slidestyle.TEMPLATE_PATH) as archive:
        theme = archive.read("ppt/theme/theme1.xml").decode()
    assert 'typeface="Aptos Display"' in theme
    assert 'typeface="Aptos"' in theme
    assert 'val="1973EB"' in theme


def test_content_slide_has_title_dash_and_number(tmp_path):
    deck = slidestyle.new_deck()
    slide = slidestyle.content_slide(deck, "Measuring Galaxy Redshift")
    assert slide.shapes.title.text_frame.text == "Measuring Galaxy Redshift"
    kinds = {shape.placeholder_format.type for shape in slide.placeholders}
    assert pptx.enum.shapes.PP_PLACEHOLDER.SLIDE_NUMBER in kinds
    names = [shape.name for shape in slide.slide_layout.shapes]
    assert "Accent Dash" in names
    path = tmp_path / "deck.pptx"
    slidestyle.save(deck, path)
    assert _codes(checker.check_deck(path)) == set()


def test_markup_runs():
    runs = slidestyle._parse(
        r"**bold** *it* !!red!! ==blue== [x](https://a.b) \* $\gamma$"
    )
    styled = [(r.text, r.bold, r.italic, r.color, r.link, r.math) for r in runs]
    assert ("bold", True, False, None, None, None) in styled
    assert ("it", False, True, None, None, None) in styled
    assert ("red", False, False, slidestyle.RED, None, None) in styled
    assert ("blue", False, False, slidestyle.ACCENT, None, None) in styled
    assert ("x", False, False, None, "https://a.b", None) in styled
    assert ("*", False, False, None, None, None) in styled
    assert ("", False, False, None, None, r"\gamma") in styled


def test_bengali_runs_are_tagged():
    deck = slidestyle.new_deck()
    slide = slidestyle.blank_slide(deck)
    shape = slidestyle.text(slide, "Biprateep Dey | বিপ্রতীপ দে", 1, 1, 6)
    langs = [
        run.get_or_add_rPr().get("lang")
        for run in shape._element.iter(ns.qn("a:r"))
    ]
    assert "bn-BD" in langs


def test_figure_contain_keeps_aspect(png):
    deck = slidestyle.new_deck()
    slide = slidestyle.content_slide(deck, "A Figure")
    picture = slidestyle.figure(slide, png, (1, 1, 10, 10), alt="blank")
    assert picture.width / picture.height == pytest.approx(2.0, rel=1e-3)
    assert (
        picture._element.find(f".//{ns.qn('p:cNvPr')}").get("descr") == "blank"
    )


def test_title_case():
    assert (
        slidestyle.title_case("why do we need maps of our universe?")
        == "Why Do We Need Maps of Our Universe?"
    )
    assert slidestyle.title_case("model.fit() does nothing") == (
        "model.fit() Does Nothing"
    )


def test_appear_targets_shapes(png):
    deck = slidestyle.new_deck()
    slide = slidestyle.content_slide(deck, "Builds")
    label = slidestyle.text(slide, "Label", 1, 2, 3)
    pointer = slidestyle.arrow(slide, (1, 3), (3, 4))
    slidestyle.appear(slide, [[label, pointer]])
    targets = [
        int(element.get("spid"))
        for element in slide._element.iter(ns.qn("p:spTgt"))
    ]
    assert targets == [label.shape_id, pointer.shape_id]
    builds = [b.get("spid") for b in slide._element.iter(ns.qn("p:bldP"))]
    assert builds == [str(label.shape_id)]


@pytest.mark.skipif(
    not (_HAS_PANDOC and _HAS_LATEX), reason="needs pandoc and LaTeX"
)
def test_equation_is_native_with_picture_fallback(tmp_path):
    deck = slidestyle.new_deck()
    slide = slidestyle.content_slide(deck, "An Equation")
    wrapper = slidestyle.equation(slide, r"\frac{a}{b} + \hat{F}(z)", 2, 2)
    assert etree.QName(wrapper).namespace == _MC
    choice = wrapper.find(f"{{{_MC}}}Choice")
    assert choice.get("Requires") == "a14"
    math = choice.find(f".//{{{_A14}}}m/{{{_M}}}oMathPara")
    assert math is not None
    runs = list(math.iter(f"{{{_M}}}r"))
    assert runs
    for run in runs:
        latin = run.find(f".//{ns.qn('a:latin')}")
        assert latin.get("typeface") == "Cambria Math"
    fallback = wrapper.find(f"{{{_MC}}}Fallback")
    blip = fallback.find(f".//{ns.qn('a:blip')}")
    assert blip.get(ns.qn("r:embed")) in slide.part.rels
    path = tmp_path / "eq.pptx"
    slidestyle.save(deck, path)
    assert "font" not in _codes(checker.check_deck(path))


@pytest.mark.skipif(not _HAS_PANDOC, reason="needs pandoc")
def test_inline_math_has_text_fallback():
    deck = slidestyle.new_deck()
    slide = slidestyle.content_slide(deck, "Inline Math")
    wrapper = slidestyle.text(slide, r"probability $\gamma$", 1, 2, 6)
    fallback_text = "".join(
        t.text or ""
        for t in wrapper.find(f"{{{_MC}}}Fallback").iter(ns.qn("a:t"))
    )
    assert "γ" in fallback_text
    assert not list(wrapper.find(f"{{{_MC}}}Fallback").iter(f"{{{_A14}}}m"))


def test_checker_flags_violations(tmp_path, png):
    deck = slidestyle.new_deck()
    slide = slidestyle.content_slide(deck, "Results.")
    slidestyle.text(slide, "A 14 pt label", 1, 2, 4, size=14)
    shape = slidestyle.text(slide, "In Open Sans", 1, 3, 4)
    shape.text_frame.paragraphs[0].runs[0].font.name = "Open Sans"
    slidestyle.bullets(slide, ["one", "two"], 1, 4, 6)
    slide.notes_slide.notes_text_frame.text = "Say this."
    slidestyle.figure(slide, png, alt="blank")
    path = tmp_path / "bad.pptx"
    slidestyle.save(deck, path)
    codes = _codes(checker.check_deck(path))
    assert {"title-period", "text-size", "font", "bullets", "notes"} <= codes
    assert "notes" not in _codes(checker.check_deck(path, allow_notes=True))


@pytest.mark.filterwarnings("ignore:skipping asset")
def test_checker_allows_house_exceptions(tmp_path, png):
    deck = slidestyle.new_deck()
    slidestyle.title_slide(deck, "**Lead:** A Talk")
    slidestyle.statement_slide(deck, "Why does this matter?")
    slide = slidestyle.content_slide(deck, "A Figure Slide")
    slidestyle.figure(slide, png, alt="blank")
    slidestyle.credit(slide, "Image courtesy: Someone")
    slidestyle.summary_slide(deck, ["First **point**.", "Second **point**."])
    slidestyle.divider_slide(deck)
    path = tmp_path / "good.pptx"
    slidestyle.save(deck, path)
    assert _codes(checker.check_deck(path)) == set()


def test_checker_title_drift(tmp_path, png):
    deck = slidestyle.new_deck()
    for title in ("First Slide", "Second Slide", "Third Slide"):
        slide = slidestyle.content_slide(deck, title)
        slidestyle.figure(slide, png, alt="blank")
    slide.shapes.title.top = pptx.util.Inches(2)
    path = tmp_path / "drift.pptx"
    slidestyle.save(deck, path)
    drifted = [
        f.slide for f in checker.check_deck(path) if f.code == "title-drift"
    ]
    assert drifted == [3]


@pytest.mark.skipif(
    not (shutil.which("latex") and shutil.which("dvisvgm")),
    reason="needs latex",
)
def test_latex_to_svg_writes_outlines(tmp_path):
    output = tmp_path / "eq.svg"
    width, height = svg_tool.render(r"\frac{a}{b}", output, 28)
    svg = output.read_text()
    assert "<text" not in svg
    assert width > 0
    assert height > 0


def test_check_slides_cli_exit_status(tmp_path, png):
    deck = slidestyle.new_deck()
    slide = slidestyle.content_slide(deck, "Bold Title")
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.bold = True
    slidestyle.figure(slide, png, alt="blank")
    path = tmp_path / "bold.pptx"
    slidestyle.save(deck, path)
    result = subprocess.run(
        [sys.executable, str(_ROOT / "scripts" / "check_slides.py"), str(path)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 1
    assert "title-bold" in result.stdout
